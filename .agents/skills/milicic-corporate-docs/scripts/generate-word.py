#!/usr/bin/env python3
"""
generate-word.py - Generador de Documentos Word Institucionales de Milicic S.A.
Genera archivos .docx con tipografía, paleta de colores corporativa (Naranja Milicic y Slate Dark),
portada ejecutiva, tablas estilizadas y bloques de callout.
Invocación recomendada: uv run --with python-docx python generate-word.py [output.docx]
"""

import sys
import os

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Aplica color de fondo a una celda de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Establece márgenes internos (padding) para una celda en twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def build_milicic_word_doc(output_path="informe-tecnico-milicic.docx"):
    doc = Document()

    # Configuración de márgenes estándar A4
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)   # A4
        section.page_height = Inches(11.69)

    # Colores corporativos Milicic
    COLOR_PRIMARY = RGBColor(234, 88, 12)     # #EA580C Naranja Milicic
    COLOR_SLATE_DARK = RGBColor(15, 23, 42)   # #0F172A
    COLOR_SLATE_LEAD = RGBColor(30, 41, 59)   # #1E293B
    COLOR_SLATE_BODY = RGBColor(51, 65, 85)   # #334155
    COLOR_MUTED = RGBColor(100, 116, 139)     # #64748B

    HEX_PRIMARY = "EA580C"
    HEX_SLATE_DARK = "0F172A"
    HEX_LIGHT_ORANGE = "FFF7ED"
    HEX_ZEBRA = "F8FAFC"
    HEX_BORDER = "E2E8F0"

    # Configurar estilo Normal
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = COLOR_SLATE_BODY

    # ==================== PORTADA EJECUTIVA ====================
    
    # 1. Píldora de Categoría
    p_cat = doc.add_paragraph()
    p_cat.paragraph_format.space_before = Pt(36)
    p_cat.paragraph_format.space_after = Pt(12)
    r_cat = p_cat.add_run("INFORME EJECUTIVO DE ARQUITECTURA & TI")
    r_cat.font.name = 'Segoe UI'
    r_cat.font.size = Pt(9)
    r_cat.font.bold = True
    r_cat.font.color.rgb = COLOR_PRIMARY

    # 2. Título Principal
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(10)
    r_title = p_title.add_run("Infraestructura de Autenticación Distribuida\nen Oficinas Satélites")
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_SLATE_DARK

    # 3. Subtítulo
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(28)
    r_sub = p_sub.add_run("Continuidad operativa, modo isla y resiliencia en campamentos mineros ante contingencias de enlace WAN.")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(11.5)
    r_sub.font.color.rgb = COLOR_MUTED

    # 4. Tabla de Metadatos
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta_table.autofit = False

    metadata = [
        ("Organización:", "Milicic S.A. | Minería, Construcción e Infraestructura"),
        ("Área Responsable:", "Gerencia de Infraestructura y Tecnología de la Información"),
        ("Fecha de Publicación:", "Septiembre 2026  •  Versión 1.0 Final"),
        ("Clasificación:", "Confidencial - Uso Interno Exclusivo")
    ]

    for idx, (label, val) in enumerate(metadata):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(1.8)
        cell_val.width = Inches(4.5)

        p_lbl = cell_lbl.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9)
        r_lbl.font.color.rgb = COLOR_SLATE_LEAD

        p_v = cell_val.paragraphs[0]
        r_v = p_v.add_run(val)
        r_v.font.size = Pt(9)
        r_v.font.color.rgb = COLOR_SLATE_BODY

        set_cell_margins(cell_lbl, top=60, bottom=60, left=60, right=100)
        set_cell_margins(cell_val, top=60, bottom=60, left=100, right=60)

    # Salto de página tras la portada
    doc.add_page_break()

    # ==================== SECCIÓN 1: RESUMEN EJECUTIVO ====================
    
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(18)
    p_h1.paragraph_format.space_after = Pt(8)
    r_h1 = p_h1.add_run("1. Resumen Ejecutivo")
    r_h1.font.name = 'Segoe UI'
    r_h1.font.size = Pt(15)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_PRIMARY

    p_body1 = doc.add_paragraph(
        "El presente proyecto implementa una arquitectura de autenticación local resiliente para los campamentos "
        "mineros y obradores de Milicic S.A. desplegados a lo largo del país. La solución desacopla la autenticación "
        "del personal técnico y operativo de la disponibilidad constante de la troncal WAN hacia el Datacenter Central."
    )
    p_body1.paragraph_format.space_after = Pt(12)

    # Callout Box (Tabla de 1x1 con borde izquierdo naranja)
    callout_table = doc.add_table(rows=1, cols=1)
    callout_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout_cell = callout_table.rows[0].cells[0]
    callout_cell.width = Inches(6.27)
    set_cell_background(callout_cell, HEX_LIGHT_ORANGE)
    set_cell_margins(callout_cell, top=140, bottom=140, left=200, right=200)

    # Borde izquierdo naranja en XML
    tcPr = callout_cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                          f'<w:top w:val="none"/>'
                          f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_PRIMARY}"/>'
                          f'<w:bottom w:val="none"/>'
                          f'<w:right w:val="none"/>'
                          f'</w:tcBorders>')
    tcPr.append(tcBorders)

    p_call = callout_cell.paragraphs[0]
    r_call_title = p_call.add_run("Principio de Continuidad Operativa (Zero Lockout):\n")
    r_call_title.font.bold = True
    r_call_title.font.size = Pt(9.5)
    r_call_title.font.color.rgb = COLOR_PRIMARY

    r_call_desc = p_call.add_run(
        "Ante la caída imprevista del enlace satelital Starlink por condiciones meteorológicas extremas "
        "o tormentas geomagnéticas, los usuarios continúan autenticándose contra la réplica de Samba AD DC "
        "y FreeRADIUS alojada localmente en faena, sin interrupción de turnos."
    )
    r_call_desc.font.size = Pt(9)
    r_call_desc.font.color.rgb = COLOR_SLATE_BODY

    # ==================== SECCIÓN 2: MATRIZ DE SITIOS ====================
    
    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(20)
    p_h2.paragraph_format.space_after = Pt(8)
    r_h2 = p_h2.add_run("2. Estado de Infraestructura por Sitio")
    r_h2.font.name = 'Segoe UI'
    r_h2.font.size = Pt(13)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_SLATE_DARK

    table_data = [
        ("Sitio / Campamento", "Ubicación", "Enlace Primario", "Hardware", "Estado"),
        ("Campamento Los Azules", "San Juan (Cordillera)", "Starlink Business", "Raspberry Pi 5 8GB", "OPERATIVO"),
        ("Litio Salar Hombre Muerto", "Catamarca (Puna)", "Starlink + 4G Backup", "Raspberry Pi 5 8GB", "OPERATIVO"),
        ("Obrador Autopista RN 34", "Santa Fe", "Fibra Óptica 100M", "Micro-Server Linux", "OPERATIVO"),
        ("Parque Eólico La Angostura", "Chubut (Patagonia)", "Starlink Standard", "Raspberry Pi 4 8GB", "EN OBSERVACIÓN"),
    ]

    doc_table = doc.add_table(rows=len(table_data), cols=5)
    doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    doc_table.autofit = False

    col_widths = [Inches(1.8), Inches(1.3), Inches(1.2), Inches(1.2), Inches(0.9)]

    for row_idx, row_values in enumerate(table_data):
        row = doc_table.rows[row_idx]
        is_header = (row_idx == 0)
        is_even = (row_idx % 2 == 0)

        for col_idx, val in enumerate(row_values):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            run = p.add_run(val)
            run.font.name = 'Segoe UI'

            if is_header:
                set_cell_background(cell, HEX_SLATE_DARK)
                set_cell_margins(cell, top=140, bottom=140, left=120, right=120)
                run.font.bold = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(255, 255, 255)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [3, 4] else WD_ALIGN_PARAGRAPH.LEFT
            else:
                bg = HEX_ZEBRA if is_even else "FFFFFF"
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                run.font.size = Pt(8.5)
                if col_idx == 4:
                    run.font.bold = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run.font.color.rgb = RGBColor(22, 163, 74) if "OPERATIVO" in val else RGBColor(217, 119, 6)
                else:
                    run.font.color.rgb = COLOR_SLATE_BODY
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Guardar documento
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    print(f"✅ Documento Word Milicic generado con éxito: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "informe-tecnico-milicic.docx"
    build_milicic_word_doc(out)
